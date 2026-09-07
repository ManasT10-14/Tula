"""Primary-field evidence, uncertainty and correction history regressions.

Text is deliberately injected OCR fixture data, not a photograph accuracy claim.
"""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape
from PIL import Image, ImageDraw
from pypdf import PdfReader

from tula.domain.enums import DeclarationClass as DC
from tula.domain.enums import ExtractionPath, Panel
from tula.domain.models import Analysis, Declaration, PackageFacts, Scan, sha256_file
from tula.extract.pipeline import extract
from tula.extract.provenance import get_provenance
from tula.ocr.base import OcrLine, OcrResult
from tula.report import docx_export, evidence, pdf, render
from tula.rules.engine import RulesEngine
from tula.security import User
from tula.services import review
from tula.storage.db import Repository

OWNER = User("provenance-officer", "provenance", "Provenance QA Inspector", "inspector", True)
REASON = "Read the selected source rectangle and verified the printed marking."


def line(text, box=(20, 20, 260, 50), confidence=.98, frame="front.png"):
    return OcrLine(text, box, confidence=confidence, frame=frame)


def extract_lines(*lines):
    return extract([(Panel.PDP, OcrResult(lines=list(lines), width=900, height=450))])


def test_cue_only_has_provenance_without_becoming_normalized():
    result = extract_lines(line("MRP", confidence=.91))
    declaration = result.declarations[DC.RETAIL_SALE_PRICE]
    assert declaration.norm == {}  # Rules distinguish an unread cue from a value.
    provenance = declaration.provenance
    assert provenance.status == "cue_only" and provenance.ocr_confidence == .91
    assert provenance.extraction_confidence is None
    assert provenance.sources[0].bbox == (20, 20, 260, 50)
    assert provenance.sources[0].frame == "front.png"
    assert result.intelligence["declarations"][DC.RETAIL_SALE_PRICE.value]["status"] == "cue_only"


def test_low_confidence_address_controls_entire_contact_block_score():
    result = extract_lines(line("Manufactured by Example Ltd", confidence=.99),
                           line("Plot 22, Pune 411018", (20, 57, 280, 87), confidence=.56))
    declaration = result.declarations[DC.MANUFACTURER]
    provenance = declaration.provenance
    assert declaration.norm["has_address"]
    assert declaration.bbox == (20, 20, 260, 50)  # Glyph metrology retains the anchor.
    assert declaration.confidence == provenance.ocr_confidence == .56
    assert provenance.extraction_confidence == pytest.approx(.504)
    assert provenance.method == "contact_block"
    assert [source.text for source in provenance.sources] == declaration.raw.splitlines()


def test_conflicting_prices_keep_every_frame_and_conservative_score():
    result = extract([(Panel.PDP, OcrResult(lines=[line("MRP Rs 100", confidence=.98)], width=900, height=450)),
                      (Panel.BACK, OcrResult(lines=[line("MRP Rs 200 inclusive of all taxes", confidence=.61,
                                                         frame="back.png")], width=900, height=450))])
    declaration = result.declarations[DC.RETAIL_SALE_PRICE]
    assert declaration.norm["count"] == 2
    assert declaration.provenance.status == "needs_review"
    assert declaration.confidence == .61
    assert {source.frame for source in declaration.provenance.sources} == {"front.png", "back.png"}
    assert len(declaration.norm["source_spans"]) == 2


def test_ambiguous_spatial_values_keep_all_candidates_without_inventing_price():
    result = extract_lines(line("MRP", (20, 20, 75, 40)),
                           line("120", (90, 20, 145, 40), .92),
                           line("180", (20, 52, 75, 72), .87))
    declaration = result.declarations[DC.RETAIL_SALE_PRICE]
    assert declaration.norm["requires_review"] and declaration.norm.get("value") is None
    assert declaration.provenance.status == "needs_review"
    assert {source.text for source in declaration.provenance.sources} == {"MRP", "120", "180"}
    assert declaration.provenance.ocr_confidence == .87


def test_generic_and_brand_methods_describe_actual_heuristics():
    declarations = extract_lines(line("ACME", (20, 20, 170, 65)),
                                 line("Tea", (20, 100, 170, 130))).declarations
    assert declarations[DC.GENERIC_NAME].provenance.method == "generic_name_lexicon"
    assert declarations[DC.BRAND].provenance.method == "brand_layout_heuristic"


def test_unlocated_historical_field_does_not_invent_confidence():
    declaration = Declaration(klass=DC.BRAND, raw="Legacy brand", norm={"name": "Legacy brand"})
    before = declaration.model_dump_json()
    provenance = get_provenance(declaration)
    assert provenance.ocr_confidence is None and provenance.extraction_confidence is None
    assert provenance.method == "not_recorded" and provenance.sources == []
    assert declaration.model_dump_json() == before


def make_case(directory):
    directory.mkdir(parents=True, exist_ok=True)
    frames = []
    for name, texts in [("front", ["PROVENANCE QA FIXTURE", "MRP Rs 100", "Manufactured by Example Ltd", "Plot 22, Pune 411018"]),
                        ("back", ["PROVENANCE QA FIXTURE", "MRP Rs 200 inclusive of all taxes"] )]:
        frame = directory / f"{name}.png"
        image = Image.new("RGB", (900, 450), "white")
        draw = ImageDraw.Draw(image)
        for y, text in enumerate(texts):
            draw.text((20, 15 + y * 65), text, fill="black", font_size=26)
        image.save(frame)
        frames.append(str(frame))
    result = extract([(Panel.PDP, OcrResult(lines=[
        line("MRP Rs 100", (20, 75, 250, 110), .98, frames[0]),
        line("Manufactured by Example Ltd", (20, 140, 600, 175), .99, frames[0]),
        line("Plot 22, Pune 411018", (20, 205, 430, 240), .56, frames[0]),
    ], width=900, height=450)), (Panel.BACK, OcrResult(lines=[
        line("MRP Rs 200 inclusive of all taxes", (20, 75, 680, 110), .61, frames[1])
    ], width=900, height=450))])
    rules = RulesEngine.from_directory()
    analysis = Analysis(scan=Scan(scan_id="PROVENANCE-QA-001", frames=frames,
        frame_hashes={frame: sha256_file(frame) for frame in frames},
        source="bench", inspector_id=OWNER.id, operator=OWNER.display_name), package=PackageFacts(),
        declarations=result.declarations, spans=result.spans, intelligence=result.intelligence,
        rules_version=rules.pack.version, engine="injected-ocr-qa-fixture",
        warnings=["Generated source images and injected OCR are provenance QA fixtures, not a real inspection."])
    repo = Repository(directory / "provenance.db")
    repo.archive_rules(rules.pack)
    repo.save(analysis)
    return SimpleNamespace(a=analysis, repo=repo, rules=rules)


@pytest.fixture
def case(tmp_path):
    return make_case(tmp_path)


def correct(case, raw="MRP Rs 200 inclusive of all taxes", revision=0, kind=DC.RETAIL_SALE_PRICE):
    return review.correct(case.repo, case.rules, case.a.scan.scan_id, revision, OWNER,
                          kind, raw, 1, (20, 75, 680, 110), REASON)


def revision_zero_bytes(case):
    with case.repo._connect() as connection:
        return connection.execute("SELECT record FROM inspection_revision WHERE scan_id=? AND revision=0",
                                  (case.a.scan.scan_id,)).fetchone()["record"]


def test_repeated_correction_preserves_original_ocr_and_immutable_revision(case):
    original = case.a.declarations[DC.RETAIL_SALE_PRICE]
    initial_bytes = revision_zero_bytes(case)
    first = correct(case)
    second = correct(case, "MRP Rs 210 inclusive of all taxes", first.review.revision)
    field = second.declarations[DC.RETAIL_SALE_PRICE]
    provenance = field.provenance
    assert field.path is ExtractionPath.MANUAL
    assert provenance.ocr_confidence is None and provenance.extraction_confidence is None
    assert provenance.original_transcription == original.raw
    assert provenance.original_sources == original.provenance.sources
    assert provenance.original_ocr_confidence == .61
    assert provenance.sources[0].ocr_confidence is None
    assert provenance.sources[0].frame == case.a.scan.frames[1]
    assert provenance.sources[0].image_width == 900
    assert field.norm["original_ocr"] == original.raw
    assert second.review.corrections[-1]["before"]["raw"] == first.declarations[DC.RETAIL_SALE_PRICE].raw
    assert second.review.corrections[-1]["after"]["provenance"]["original_transcription"] == original.raw
    assert second.intelligence["declarations"][DC.RETAIL_SALE_PRICE.value]["ocr_confidence"] is None
    assert revision_zero_bytes(case) == initial_bytes


def test_added_field_never_becomes_original_ocr_after_second_correction(case):
    assert DC.GENERIC_NAME not in case.a.declarations
    first = correct(case, "QA commodity", kind=DC.GENERIC_NAME)
    second = correct(case, "Rechecked QA commodity", first.review.revision, DC.GENERIC_NAME)
    provenance = second.declarations[DC.GENERIC_NAME].provenance
    assert provenance.original_transcription is None and provenance.original_sources == []
    assert provenance.original_ocr_confidence is None and provenance.ocr_confidence is None


def test_legacy_correction_recovers_revision_zero_not_previous_manual_text(case):
    first = correct(case)
    original = case.a.declarations[DC.RETAIL_SALE_PRICE]
    def simulate_legacy_record(analysis):
        field = analysis.declarations[DC.RETAIL_SALE_PRICE]
        field.provenance = None
        field.norm["original_ocr"] = "A prior manually corrected value"
    legacy = case.repo.revise(first.scan.scan_id, first.review.revision, OWNER.id,
                              "qa.legacy_fixture", "Simulate historical correction schema", simulate_legacy_record)
    second = correct(case, "MRP Rs 210 inclusive of all taxes", legacy.review.revision)
    assert second.declarations[DC.RETAIL_SALE_PRICE].provenance.original_transcription == original.raw
    assert second.declarations[DC.RETAIL_SALE_PRICE].provenance.original_sources == original.provenance.sources


def render_inspection(analysis):
    env = Environment(loader=FileSystemLoader(Path(__file__).parents[1] / "src/tula/web/templates"),
                      autoescape=select_autoescape(["html"]))
    request = SimpleNamespace(state=SimpleNamespace(user=OWNER, csrf_token="test-csrf"),
                              url=SimpleNamespace(path=f"/inspections/{analysis.scan.scan_id}"))
    return env.get_template("inspection.html").render(a=analysis, request=request, sections=render.build(analysis),
        integrity=evidence.verify(analysis), headline=render.headline(analysis), history=[], shrinkflation=[],
        declaration_labels=render.DECLARATION_LABEL, finding_rows=lambda finding: render.finding_rows(finding, analysis),
        finding_message=render.finding_message,
        verdict_class=lambda _: "warn", verdict_label=render.VERDICT_LABEL, csrf_token="test-csrf",
        asset_url=lambda path: f"/static/{path}")


def test_inspection_shows_all_sources_and_corrected_original_text_safely(case):
    updated = correct(case, "MRP Rs 200 <script>alert(1)</script> inclusive of all taxes")
    html = render_inspection(updated)
    assert "View field evidence and confidence" in html
    assert "Open field evidence image 2" in html and "Open original OCR source image 1" in html
    assert 'href="/inspections/PROVENANCE-QA-001/frames/1"' in html and 'viewBox="' in html
    assert "Original OCR transcription" in html and "Original OCR model score: 0.61" in html
    assert "no automated OCR or extraction score is assigned to the corrected text" in html
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "OCR confidence 61%" not in html


def test_pdf_docx_and_crops_preserve_cross_image_original_and_manual_evidence(case, tmp_path):
    updated = correct(case)
    # Keep one real recomputed finding in this bounded export fixture.
    updated.findings = [finding for finding in updated.findings if finding.declaration == DC.RETAIL_SALE_PRICE][:1]
    pdf_path = pdf.write(updated, tmp_path / "provenance.pdf")
    docx_path = docx_export.write(updated, tmp_path / "provenance.docx")
    reader = PdfReader(pdf_path)
    pdf_text = "\n".join(page.extract_text() for page in reader.pages)
    document = Document(docx_path)
    docx_text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for text in (pdf_text, docx_text):
        assert "Original OCR transcription: MRP Rs 100" in text
        assert "Original OCR model score: 0.61" in text
        assert "Original evidence: Image 1" in text and "Original evidence: Image 2" in text
        assert "officer corrected" in text and "no automated OCR or extraction score" in text
        assert case.a.scan.frames[0] not in text  # Public report text uses image references.
    attached = json.loads(reader.attachments["analysis.json"][0])
    assert attached["declarations"]["retail_sale_price"]["provenance"]["ocr_confidence"] is None
    assert {item[1] for item in evidence.crops(updated)} == set(case.a.scan.frames)
    assert all(item["status"] == "verified" for item in evidence.verify(updated))
