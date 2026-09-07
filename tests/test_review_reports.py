"""Exports must preserve review state and evidence without inventing approval."""
import json
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader

from tula.domain.enums import DeclarationClass as DC
from tula.domain.enums import Panel, Verdict
from tula.domain.models import (
    Analysis,
    Citation,
    Declaration,
    Finding,
    PackageFacts,
    ReviewDecision,
    ReviewState,
    Scan,
    sha256_file,
)
from tula.extract.intelligence import extract_intelligence
from tula.ocr.base import OcrLine
from tula.report import docx_export, evidence, pdf, render


def reviewed_analysis(directory: Path) -> Analysis:
    directory.mkdir(parents=True, exist_ok=True)
    frame = directory / "report-qa-label.png"
    image = Image.new("RGB", (900, 450), "white")
    drawer = ImageDraw.Draw(image)
    for y, text in [(25, "REPORT QA SAMPLE"), (90, "MRP Rs 120"),
                    (160, "EXP 08/09/26"), (230, "Ingredients: Milk solids, wheat flour")]:
        drawer.text((25, y), text, fill="black", font_size=28)
    image.save(frame)
    name = str(frame)
    declaration = Declaration(klass=DC.RETAIL_SALE_PRICE, raw="MRP Rs 120", norm={"value": 120, "currency": "INR"},
                              bbox=(20, 85, 350, 125), frame=name, panel=Panel.PDP, confidence=.94)
    findings = [Finding(finding_id="QA-MRP", rule_id="LMPCR.R6.MRP", rules_version="draft-qa",
                        citation=Citation(clause="Rule 6", text="Draft screening criterion"),
                        declaration=DC.RETAIL_SALE_PRICE, verdict=Verdict.VIOLATION,
                        message="Tax wording requires verification", extracted="MRP Rs 120"),
                Finding(finding_id="QA-DATE", rule_id="LMPCR.R6.DATE", rules_version="draft-qa",
                        citation=Citation(clause="Rule 6"), verdict=Verdict.INCONCLUSIVE,
                        message="Printed date interpretation requires review")]
    now = datetime(2026, 9, 7, 12, 0, tzinfo=UTC)
    review = ReviewState(revision=7, status="approved", submitted_by="qa-inspector",
                         approved_by="QA Supervisor", approved_at=now,
                         approval_reason="Seeded review state to verify report rendering", decisions={
        "QA-MRP": ReviewDecision(finding_id="QA-MRP", verdict=Verdict.VIOLATION,
            reason="QA fixture decision; this is not a real inspection", actor_id="qa-inspector", actor_name="QA Inspector", created_at=now),
        "QA-DATE": ReviewDecision(finding_id="QA-DATE", verdict=Verdict.PASS,
            reason="QA fixture resolution used to test export status", actor_id="qa-inspector", actor_name="QA Inspector", created_at=now),
    }, corrections=[{"kind": "retail_sale_price", "before": {"raw": "MRP Rs 180", "norm": {"value": 180}},
                      "after": declaration.model_dump(mode="json"), "actor": "QA Inspector", "actor_id": "qa-inspector",
                      "reason": "QA correction history fixture", "created_at": now.isoformat()}],
       comments=[{"actor": "QA Inspector", "action": "comment", "created_at": now.isoformat(),
                  "text": "Generated label and seeded decisions are report QA data."}])
    intelligence = extract_intelligence([
        (Panel.PDP, OcrLine("EXP 08/09/26", (20, 155, 400, 195), confidence=.92, frame=name)),
        (Panel.PDP, OcrLine("Ingredients: Milk solids, wheat flour", (20, 225, 700, 265), confidence=.95, frame=name)),
    ], allergen_concerns=["milk", "gluten"])
    return Analysis(scan=Scan(scan_id="QA-REPORT-001", frames=[name], frame_hashes={name: sha256_file(name)},
                             operator="QA Inspector", source="bench", captured_at=now),
                    package=PackageFacts(), declarations={DC.RETAIL_SALE_PRICE: declaration},
                    findings=findings, rules_version="draft-qa", engine="report-qa-fixture",
                    review=review, intelligence=intelligence,
                    warnings=["Generated label and seeded officer decisions are test data for report verification."])


@pytest.fixture
def analysis(tmp_path):
    return reviewed_analysis(tmp_path)


def test_report_summary_distinguishes_machine_and_review(analysis):
    text = "\n".join(str(section.rows) for section in render.build(analysis))
    assert "Machine potential violations" in text and "Inspector verified violations" in text
    assert "QA Supervisor" in text and "PENDING" not in text
    assert "Record revision" in text and "7" in text
    assert "supervisor review approved" in render.headline(analysis).lower()
    assert "not statutory certification" in render.scope_note(analysis)


def test_draft_report_never_claims_recorded_verification(analysis):
    analysis.review = ReviewState()
    assert "Review required" in render.headline(analysis)
    assert "SUPERVISOR REVIEW APPROVED" not in render.review_banner(analysis)
    rows = dict(render.finding_rows(analysis.findings[0], analysis))
    assert rows["Inspector decision"].startswith("Not recorded")


def test_officer_decision_does_not_overwrite_machine_verdict(analysis):
    rows = dict(render.finding_rows(analysis.findings[1], analysis))
    assert rows["Machine result"] == "Inconclusive"
    assert rows["Inspector decision"] == "PASS"
    assert rows["Officer reason"] == "QA fixture resolution used to test export status"


@pytest.mark.parametrize("verdict", [Verdict.INCONCLUSIVE, Verdict.UNVERIFIED])
def test_uncertain_headings_do_not_assert_a_stored_negative_message(analysis, tmp_path, verdict):
    finding = analysis.findings[0].model_copy(update={"verdict": verdict,
        "message": "No retail sale price is declared on the package."})
    analysis.findings[0] = finding
    original = analysis.model_dump_json()
    shown = render.finding_message(finding)
    assert "Retail sale price" in shown and finding.message not in shown
    pdf.write(analysis, tmp_path / "uncertain.pdf")
    docx_export.write(analysis, tmp_path / "uncertain.docx")
    reader = PdfReader(tmp_path / "uncertain.pdf")
    assert finding.message not in "\n".join(page.extract_text() for page in reader.pages)
    assert finding.message not in "\n".join(p.text for p in Document(tmp_path / "uncertain.docx").paragraphs)
    assert json.loads(reader.attachments["analysis.json"][0])["findings"][0]["message"] == finding.message
    assert analysis.model_dump_json() == original


def test_determined_finding_heading_retains_its_recorded_message(analysis):
    finding = analysis.findings[0]
    assert finding.verdict == Verdict.VIOLATION
    assert render.finding_message(finding) == finding.message


def test_exports_include_revision_corrections_dates_allergens_and_images(analysis, tmp_path):
    exported_pdf = pdf.write(analysis, tmp_path / "review.pdf")
    exported_docx = docx_export.write(analysis, tmp_path / "review.docx")
    reader = PdfReader(exported_pdf)
    text = "\n".join(page.extract_text() for page in reader.pages)
    for expected in ["Revision 7", "QA Supervisor", "Before correction", "MRP Rs 180", "MRP Rs 120",
                     "2026-08-09", "2026-09-08", "allergen", "Original image appendix", "Evidence crops"]:
        assert expected in text
    attached = json.loads(reader.attachments["analysis.json"][0])
    assert attached["review"]["revision"] == 7 and attached["review"]["status"] == "approved"
    assert sum(len(page.images) for page in reader.pages) >= 4
    doc = Document(exported_docx)
    assert len(doc.inline_shapes) >= 4
    assert any(p.style.name == "Title" for p in doc.paragraphs)
    with ZipFile(exported_docx) as archive:
        xml = archive.read("word/document.xml").decode()
        assert "QA Supervisor" in xml and "MRP Rs 180" in xml and "w:tblHeader" in xml
        assert "Highlighted" in xml and "PENDING REVIEW" not in xml


def test_evidence_crops_highlight_source_without_changing_original(analysis):
    frame = analysis.scan.frames[0]
    before = sha256_file(frame)
    crops = list(evidence.crops(analysis))
    assert len(crops) == 3
    with Image.open(crops[0][3]) as im:
        assert (181, 54, 19) in im.getdata()
    assert sha256_file(frame) == before


def test_changed_evidence_is_not_embedded_in_reports(analysis):
    Image.new("RGB", (900, 450), "black").save(analysis.scan.frames[0])
    assert not list(evidence.crops(analysis))
    assert not list(evidence.originals(analysis))
    assert evidence.verify(analysis)[0]["status"] == "mismatch"


def test_notice_uses_only_verified_findings(analysis, tmp_path):
    analysis.review = ReviewState()
    exported = docx_export.write_notice(analysis, tmp_path / "notice.docx")
    text = "\n".join(p.text for p in Document(exported).paragraphs)
    assert "No inspector-verified violation" in text
    assert "Tax wording requires verification" not in text
    assert "not a served notice" in text


def test_long_review_reason_exports_without_layout_error(analysis, tmp_path):
    decision = analysis.review.decisions["QA-MRP"].model_copy(update={"reason": "Detailed officer review observation. " * 300})
    analysis.review.decisions["QA-MRP"] = decision
    path = pdf.write(analysis, tmp_path / "long-review.pdf")
    text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    # A page footer may interrupt the extracted sentence; all 300 starts and
    # ends must still be present across split table rows.
    assert text.count("Detailed") == 300
    assert text.count("observation.") == 300


@pytest.mark.parametrize(("context", "expected"), [
    ({}, "Not confirmed"),
    ({"is_imported": False}, "Not confirmed"),
    ({"is_imported": True}, "Not confirmed"),
    ({"imported_confirmed": True}, "Not confirmed"),
    ({"imported_confirmed": True, "is_imported": "false"}, "Not confirmed"),
    ({"imported_confirmed": "true", "is_imported": False}, "Not confirmed"),
    ({"imported_confirmed": True, "is_imported": False}, "No (officer-confirmed)"),
    ({"imported_confirmed": True, "is_imported": True}, "Yes (officer-confirmed)"),
])
def test_origin_export_requires_typed_officer_confirmation(analysis, context, expected):
    # The old boolean field can contain a machine inference and cannot confirm origin.
    analysis.package.is_imported = not context.get("is_imported", False)
    analysis.package.legal_context = context
    section = next(s for s in render.build(analysis) if s.title == "Applicability and exemptions")
    assert dict(section.rows)["Imported commodity"] == expected


@pytest.mark.parametrize(("context", "expected"), [
    ({"category": "food"}, "Not confirmed"),
    ({"category_confirmed": True, "category": "unknown"}, "Not confirmed"),
    ({"category_confirmed": True, "category": "invented"}, "Not confirmed"),
    ({"category_confirmed": "true", "category": "food"}, "Not confirmed"),
    ({"category_confirmed": True, "category": "food"}, "Food (officer-confirmed)"),
])
def test_suggested_category_is_separate_from_confirmed_context(analysis, context, expected):
    analysis.package.legal_context = context
    analysis.intelligence["product"] = {"category": "household"}
    sections = render.build(analysis)
    applicability = next(s for s in sections if s.title == "Applicability and exemptions")
    assert dict(applicability.rows)["Product category"] == expected
    suggestions = next(s for s in sections if s.title == "Product context observations")
    assert suggestions.rows[0][1] == "household"
    assert "Officer confirmation required" in suggestions.rows[0][2]


def observation_analysis(directory: Path) -> Analysis:
    analysis = reviewed_analysis(directory)
    frame = analysis.scan.frames[0]
    with Image.open(frame) as previous:
        image = Image.new("RGB", (900, 700), "white")
        image.paste(previous, (0, 0))
    drawer = ImageDraw.Draw(image)
    lines = [
        OcrLine("Milk biscuits", (20, 325, 350, 365), confidence=.96, frame=frame),
        OcrLine("Vegetarian", (20, 395, 350, 435), confidence=.49, frame=frame),
        OcrLine("Ingredients: stabiliser INS 322", (20, 465, 750, 505), confidence=.91, frame=frame),
        OcrLine("Made in India", (20, 535, 400, 575), confidence=.93, frame=frame),
        OcrLine("Bottle", (20, 605, 350, 645), confidence=.92, frame=frame),
    ]
    for line in lines:
        drawer.text((25, line.bbox[1] + 5), line.text, fill="black", font_size=28)
    image.save(frame)
    analysis.scan.frame_hashes[frame] = sha256_file(frame)
    analysis.scan.region = "Bengaluru Urban"
    analysis.scan.geo = (12.9716, 77.5946)
    analysis.package.legal_context = {"category": "food", "category_confirmed": True,
                                      "attested_by": "qa-inspector"}
    extra = extract_intelligence([(Panel.PDP, line) for line in lines])
    for key in ("product", "dietary", "additives"):
        analysis.intelligence[key] = extra[key]
    return analysis


def test_pdf_and_docx_export_region_and_located_uncertain_observations(tmp_path):
    analysis = observation_analysis(tmp_path)
    before = analysis.model_dump(mode="json")
    exported_pdf = pdf.write(analysis, tmp_path / "observations.pdf")
    exported_docx = docx_export.write(analysis, tmp_path / "observations.docx")
    pdf_text = " ".join("\n".join(page.extract_text() for page in PdfReader(exported_pdf).pages).split())
    doc = Document(exported_docx)
    docx_text = " ".join("\n".join([p.text for p in doc.paragraphs] + [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]).split())
    for text in (pdf_text, docx_text):
        for value in ("Bengaluru Urban; GPS 12.97160, 77.59460", "Not confirmed",
                      "Food (officer-confirmed)", "Suggested category", "Origin signal", "domestic",
                      "Milk biscuits", "Vegetarian", "INS 322", "Printed package reference", "bottle",
                      "Status: needs review", "source OCR score 0.49", "Image 1; panel pdp",
                      "pixels 20, 395, 350, 435", "not independently certified"):
            assert value in text
    assert analysis.model_dump(mode="json") == before
    assert len(list(evidence.crops(analysis))) == 8


@pytest.mark.parametrize(("region", "geo", "expected"), [
    ("Bengaluru Urban", None, "Bengaluru Urban"),
    ("", (0, 0), "GPS 0.00000, 0.00000"),
    ("  ", None, "not recorded"),
])
def test_report_location_preserves_text_without_requiring_gps(analysis, region, geo, expected):
    analysis.scan.region, analysis.scan.geo = region, geo
    particulars = next(s for s in render.build(analysis) if s.title == "Inspection particulars")
    assert dict(particulars.rows)["Location"] == expected


def test_observations_without_sources_do_not_invent_provenance(analysis):
    analysis.intelligence = {"dietary": [{"raw": "Vegan", "status": "needs_review"}],
                             "additives": [{"code": "INS 322", "sources": [{"frame": "missing.png", "image_index": 0}]}]}
    section = next(s for s in render.build(analysis) if s.title == "Printed claims and additive identifiers")
    details = "\n".join(row[2] for row in section.rows)
    assert "No source location recorded" in details
    assert "Source image not resolved" in details
    assert "OCR model score" not in details and "Image 1" not in details
    assert "missing.png" not in details


def test_additive_provenance_selects_the_code_line_from_ingredient_block(tmp_path):
    analysis = observation_analysis(tmp_path)
    item = analysis.intelligence["additives"][0]
    # Historical ingredient blocks may retain neighboring origin/package spans.
    # Seed that recorded state explicitly so this test also survives improvements
    # to extraction's block boundaries.
    item["sources"] = [source for source in item["sources"] if "INS 322" in source["text"]]
    product = analysis.intelligence["product"]
    item["sources"].extend(product["origin_evidence"][0]["sources"])
    item["sources"].extend(product["package_type_evidence"][0]["sources"])
    before = analysis.model_dump(mode="json")
    section = next(s for s in render.build(analysis) if s.title == "Printed claims and additive identifiers")
    details = next(row[2] for row in section.rows if row[1] == "INS 322")
    assert "stabiliser INS 322" in details
    assert "Made in India" not in details and "Bottle" not in details
    assert analysis.model_dump(mode="json") == before
    names = [name for name, _, _, _, _ in evidence.crops(analysis)]
    assert names.count("additive identifier") == 1
    assert "origin evidence" in names and "package reference" in names
