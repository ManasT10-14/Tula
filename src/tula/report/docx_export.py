"""Editable DOCX export.

The problem statement asks for editable formats, and it is not a box-ticking
requirement: officers amend the language of a notice before serving it, and a
locked PDF simply fails that workflow. Two documents come out of here --
the report, and a pre-filled draft notice with the facts already in place.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..domain.models import Analysis
from . import render
from .evidence import crops, originals

_INK = RGBColor(0x0C, 0x17, 0x16)
_MUTED = RGBColor(0x56, 0x63, 0x5F)
_BRASS = RGBColor(0x85, 0x60, 0x19)


def _run(paragraph, text, *, bold=False, size=10, colour=_INK, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    return run


def write(analysis: Analysis, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _configure_doc(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    eyebrow = doc.add_paragraph()
    _run(eyebrow, "LEGAL METROLOGY (PACKAGED COMMODITIES) RULES, 2011",
         bold=True, size=8, colour=_BRASS)

    title = doc.add_paragraph(style="Title")
    _run(title, "Compliance Inspection Report", bold=True, size=20, colour=RGBColor(0, 0, 0))
    doc.add_paragraph(render.review_banner(analysis))

    lede = doc.add_paragraph()
    _run(lede, render.headline(analysis), size=11)

    meta = doc.add_paragraph()
    _run(meta,
         f"Reference {analysis.scan.scan_id}  ·  Rules version {analysis.rules_version}"
         f"  ·  Evidence Tier {analysis.scan.tier.value}  ·  Revision {analysis.review.revision}",
         size=8, colour=_MUTED)

    for section in render.build(analysis):
        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        _run(heading, section.title, bold=True, size=13)

        if section.kind == "kv":
            _kv(doc, section.rows)
        elif section.kind == "table":
            _grid(doc, section.columns, section.rows)
        elif section.kind == "text":
            for line in section.rows:
                bullet = doc.add_paragraph(style="List Bullet")
                _run(bullet, line, size=9, colour=_MUTED)
        elif section.kind == "findings":
            _findings(doc, section.rows, analysis)

        if section.note:
            note = doc.add_paragraph()
            _run(note, section.note, size=8.5, colour=_MUTED, italic=True)

    images = list(originals(analysis))
    if images:
        doc.add_heading("Original image appendix", level=1)
        intro = doc.add_paragraph("Preserved images provide context. Report illustrations are reduced and oriented for reading; recorded hashes identify the original files.")
        intro.paragraph_format.keep_with_next = True
        for name, frame, stream, size in images:
            caption = doc.add_paragraph(name + " | " + Path(frame).name, style="Caption")
            caption.paragraph_format.keep_with_next = True
            factor = min(6.0 / size[0], 4.0 / size[1])
            picture = doc.add_picture(stream, width=Inches(size[0]*factor), height=Inches(size[1]*factor))
            picture._inline.docPr.set("descr", name + " for inspection " + analysis.scan.scan_id)
    evidence = list(crops(analysis))
    if evidence:
        doc.add_heading("Evidence crops", level=1)
        for name, frame, box, stream, size in evidence:
            caption = doc.add_paragraph(f"{name} | {Path(frame).name} | pixels {box}", style="Caption")
            caption.paragraph_format.keep_with_next = True
            factor = min(6.0 / size[0], .95 / size[1])
            picture = doc.add_picture(stream, width=Inches(size[0]*factor), height=Inches(size[1]*factor))
            picture._inline.docPr.set("descr", f"Highlighted {name} evidence region at pixels {box}")
    _page_footer(doc, analysis)
    doc.core_properties.title = f"Inspection {analysis.scan.scan_id} revision {analysis.review.revision}"
    doc.save(str(path))
    return path


def _kv(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    _table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(4.85)
    for key, value in rows:
        cells = table.add_row().cells
        _run(cells[0].paragraphs[0], str(key), bold=True, size=9)
        _run(cells[1].paragraphs[0], str(value), size=9)
        cells[0].width = Inches(1.65)
        cells[1].width = Inches(4.85)
    _keep_table_rows(table)


def _keep_table_rows(table):
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_together = True
                p.paragraph_format.keep_with_next = False


def _grid(doc, columns, rows):
    if not rows:
        doc.add_paragraph("Nothing recorded.")
        return
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    _table_borders(table)
    table.autofit = False
    widths = [6.5 / len(columns)] * len(columns)
    if columns == ["Declaration", "As printed", "Panel", "Scripts"]:
        widths = [1.6, 3.2, .8, .9]
    elif len(columns) == 3 and columns[0] == "Field":
        widths = [1.25, 3.05, 2.2]
    elif len(columns) == 3 and columns[0] == "Frame":
        widths = [1.45, 4.05, 1]
    elif columns[0] == "Concern":
        widths = [1, 1.2, 1, 3.3]
    elif columns[0] == "Observation":
        widths = [1.2, 2.3, 3]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, name in zip(table.rows[0].cells, columns):
        _run(cell.paragraphs[0], str(name), bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            _run(cell.paragraphs[0], str(value), size=8.5)
    _keep_table_rows(table)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def _findings(doc, findings, analysis=None):
    for index, finding in enumerate(findings, start=1):
        rgb = render.VERDICT_RGB[finding.verdict]
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        _run(heading, f"{render.VERDICT_LABEL[finding.verdict].upper()}   ",
             bold=True, size=8, colour=RGBColor(*rgb))
        _run(heading, f"{index:02d}. {render.finding_message(finding)}", bold=True, size=10.5)
        _kv(doc, render.finding_rows(finding, analysis))
        doc.add_paragraph()


# ---------------------------------------------------------------------------
# Draft notice
# ---------------------------------------------------------------------------


def write_notice(analysis: Analysis, path: str | Path, *, premises: str = "") -> Path:
    """A pre-filled draft notice the officer edits and serves.

    Left deliberately incomplete in the places only a human can fill: the
    officer's designation, the premises, and the decision itself. The facts are
    filled in because those are what the system actually knows.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    violations = analysis.verified_violations

    doc = Document()
    _configure_doc(doc)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(header, "LEGAL METROLOGY INSPECTION", bold=True, size=12)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(title, "Draft Notice Preparation", bold=True, size=18, colour=RGBColor(0, 0, 0))
    doc.add_paragraph("Draft legal rules pack. This preparation document is not a served notice and contains no digital signature.")
    doc.add_paragraph(render.headline(analysis))

    doc.add_paragraph()
    _kv(doc, [
        ("Notice number", "Not allotted"),
        ("Inspection reference", analysis.scan.scan_id),
        ("Date of inspection", analysis.scan.captured_at.strftime("%d %B %Y")),
        ("Premises inspected", premises or "Not recorded"),
        ("Inspecting officer", analysis.scan.operator or "Not recorded"),
        ("Rules version applied", analysis.rules_version),
        ("Record revision", str(analysis.review.revision)),
        ("Workflow status", analysis.review.status),
        ("Approved by", analysis.review.approved_by or "Not approved"),
    ])

    doc.add_paragraph()
    body = doc.add_paragraph()
    _run(body,
         "The following findings were verified during the recorded officer review. "
         "The applicable provisions and notice particulars require legal confirmation before service:", size=11)

    for index, finding in enumerate(violations, start=1):
        item = doc.add_paragraph(style="List Number")
        _run(item, f"{finding.citation.clause} - {finding.message} ", size=10.5)
        decision = analysis.review.decisions[finding.finding_id]
        _run(item, f"Verified by {decision.actor_name}. Reason: {decision.reason} ", size=10)
        if finding.measured and finding.threshold is not None:
            _run(item,
                 f"Measured {finding.measured.render()} against a prescribed minimum "
                 f"of {finding.threshold:.2f} mm.",
                 size=10, colour=_MUTED)

    if not violations:
        none_found = doc.add_paragraph()
        _run(none_found, "No inspector-verified violation is recorded. Machine flags alone are not included as verified allegations.", size=11)

    doc.add_paragraph()
    closing = doc.add_paragraph()
    _run(closing,
         "An authorised reviewer must verify the applicable provisions, recipient, response period "
         "and available remedies before drafting any notice for service.", size=11)

    doc.add_paragraph()
    doc.add_paragraph()
    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(sign, "Officer signature not captured by this application", size=10)

    note = doc.add_paragraph()
    _run(note,
         "DRAFT - generated from inspection reference "
         f"{analysis.scan.scan_id}, revision {analysis.review.revision}. "
         "Unresolved and rejected machine findings are excluded from verified allegations. "
         "This editable document does not update the recorded inspection.",
         size=8, colour=_MUTED, italic=True)

    _page_footer(doc, analysis)
    doc.save(str(path))
    return path


def _page_footer(doc, analysis):
    for section in doc.sections:
        paragraph = section.footer.paragraphs[0]
        _run(paragraph, f"TATVA | {analysis.scan.scan_id} | Revision {analysis.review.revision} | Page ", size=8, colour=_MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)


def _configure_doc(doc):
    for name in ("Normal", "Title", "Subtitle"):
        style = doc.styles[name]
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.underline = False
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Heading 1"].font.color.rgb = _INK
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.styles["Caption"].font.color.rgb = _MUTED
    doc.styles["Caption"].font.size = Pt(8)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].paragraph_format.space_before = Pt(7)
    doc.styles["Caption"].paragraph_format.space_after = Pt(4)


def _table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + name)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D5DEDB")
        borders.append(border)
    table._tbl.tblPr.append(borders)
